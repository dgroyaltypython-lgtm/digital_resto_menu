from openpyxl import Workbook, load_workbook
from django.db import transaction

from .models import MenuCategory, MenuItem


EXPORT_HEADERS = [
    "Category",
    "Item Name",
    "Marathi Name",
    "Price",
    "Item Type",
    "Available",
    "Active",
    "Sort Order",
]


def export_menu():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Menu"

    sheet.append(EXPORT_HEADERS)

    categories = MenuCategory.objects.prefetch_related("items").order_by(
        "sort_order",
        "name",
    )

    for category in categories:
        for item in category.items.all():
            sheet.append(
                [
                    category.name,
                    item.name,
                    item.marathi_name,
                    float(item.price),
                    item.item_type,
                    "Yes" if item.available else "No",
                    "Yes" if item.active else "No",
                    item.sort_order,
                ]
            )

    return workbook


def import_menu(file):
    workbook = load_workbook(file, data_only=True)

    sheet = workbook.active

    headers = [
        cell.value
        for cell in sheet[1]
    ]

    required_headers = [
        "Category",
        "Item Name",
        "Marathi Name",
        "Price",
        "Item Type",
        "Available",
        "Active",
        "Sort Order",
    ]

    missing_headers = [
        header
        for header in required_headers
        if header not in headers
    ]

    if missing_headers:
        raise ValueError(
            "Missing columns: "
            + ", ".join(missing_headers)
        )

    column = {
        header: headers.index(header)
        for header in required_headers
    }

    imported = 0
    created = 0
    updated = 0

    with transaction.atomic():

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if not any(row):
                continue

            category_name = row[column["Category"]]

            item_name = row[column["Item Name"]]

            if not category_name or not item_name:
                continue

            category_name = str(
                category_name
            ).strip()

            item_name = str(
                item_name
            ).strip()

            category, _ = MenuCategory.objects.get_or_create(
                name=category_name
            )

            marathi_name = row[column["Marathi Name"]] or ""

            price = row[column["Price"]] or 0

            item_type = row[column["Item Type"]] or "Veg"

            available = str(
                row[column["Available"]] or "Yes"
            ).strip().lower() in (
                "yes",
                "true",
                "1",
            )

            active = str(
                row[column["Active"]] or "Yes"
            ).strip().lower() in (
                "yes",
                "true",
                "1",
            )

            sort_order = row[column["Sort Order"]] or 0

            item, was_created = MenuItem.objects.update_or_create(
                category=category,
                name=item_name,
                defaults={
                    "marathi_name": str(
                        marathi_name
                    ).strip(),

                    "price": price,

                    "item_type": str(
                        item_type
                    ).strip(),

                    "available": available,

                    "active": active,

                    "sort_order": int(
                        sort_order
                    ),
                },
            )

            imported += 1

            if was_created:
                created += 1
            else:
                updated += 1

    return {
        "imported": imported,
        "created": created,
        "updated": updated,
    }